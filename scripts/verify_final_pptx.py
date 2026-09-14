import os
import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\note\vd\vd15_slide-style-cloner\result\팜드라이버_농기계대리운전_IR투자제안서.pptx"
prs = Presentation(pptx_path)

print(f"Total Slides: {len(prs.slides)}")

unreplaced = []
keywords_to_flag = ["SaaS", "마케팅 분석", "Company Name", "박지민", "김성환", "InsightPro", "20YY", "OOO", "내용을 입력하세요"]

for idx, slide in enumerate(prs.slides):
    slide_texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    slide_texts.append(t)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    t = cell.text.strip()
                    if t:
                        slide_texts.append(t)
    
    print(f"=== SLIDE {idx+1:02d} ({len(slide_texts)} text blocks) ===")
    for t in slide_texts[:4]:
        print(f"  {t}")
        
    for t in slide_texts:
        for kw in keywords_to_flag:
            if kw in t:
                unreplaced.append((idx + 1, kw, t))

if unreplaced:
    print(f"\n[!] Flagged {len(unreplaced)} instances of template keywords:")
    for sn, kw, text in unreplaced:
        print(f"  Slide {sn:02d} [{kw}]: {text[:80]}")
else:
    print("\n[+] All template keywords successfully replaced!")
